"""Bounded document discovery. Research output is never an address assignment.

A recent article linking a PDF does not date the contents of that PDF. Keep the
link publication date separate and require document review before activation.
"""
import hashlib,io,time,urllib.parse,re
from bs4 import BeautifulSoup
from coverage import published_at
DAY=86400000

def document_links(body,base,hosts):
 soup=BeautifulSoup(body,'html.parser');result={}
 for a in soup.select('a[href]'):
  url=urllib.parse.urljoin(base,a['href']);u=urllib.parse.urlsplit(url)
  if u.scheme!='https' or u.username or u.password or u.port not in (None,443) or u.hostname not in hosts or u.fragment:continue
  if re.search(r'\.(?:pdf|xlsx)$',u.path,re.I):result[url]={'url':url,'label':a.get_text(' ',strip=True)[:300]}
 return list(result.values())

def inspect_document(body,kind):
 if body.startswith(b'%PDF-'):
  from pypdf import PdfReader
  pdf=PdfReader(io.BytesIO(body))
  if pdf.is_encrypted or not 0<len(pdf.pages)<=160:raise ValueError('Unsupported PDF or page limit')
  chars=0
  for p in pdf.pages:
   chars+=len(p.extract_text() or '')
   if chars>2_000_000:raise ValueError('Document text limit')
  return {'format':'pdf','pages':len(pdf.pages),'textCharacters':chars}
 if kind=='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet':
  import zipfile
  from openpyxl import load_workbook
  with zipfile.ZipFile(io.BytesIO(body)) as z:
   if len(z.infolist())>500 or sum(x.file_size for x in z.infolist())>30_000_000:raise ValueError('Workbook expansion limit')
  wb=load_workbook(io.BytesIO(body),read_only=True,data_only=True);cells=0;chars=0
  try:
   for sheet in wb:
    for row in sheet.iter_rows():
     cells+=len(row)
     if cells>100000:raise ValueError('Workbook cell limit')
     chars+=sum(len(str(c.value or '')) for c in row)
     if chars>2_000_000:raise ValueError('Document text limit')
   return {'format':'xlsx','sheets':len(wb.sheetnames),'cells':cells,'textCharacters':chars}
  finally:wb.close()
 raise ValueError('Not a PDF/XLSX document')

def watch_documents(source,fetch,now,deadline=float('inf')):
 """Read up to four topical articles and three linked reviewed-host documents."""
 report={'indexUrl':source['url'],'operatorId':source['operatorId'],'regionId':source['regionId'],'fetchedAt':now,'articles':[],'documents':[],'errors':[],'runtimeRecords':0}
 body,_,final=fetch(source['url']);soup=BeautifulSoup(body,'html.parser');urls=[final]
 if source.get('index'):
  urls=[]
  for a in soup.select('a[href]'):
   u=urllib.parse.urljoin(final,a['href']);p=urllib.parse.urlsplit(u);label=a.get_text(' ',strip=True)
   if p.scheme=='https' and p.netloc==urllib.parse.urlsplit(final).netloc and not p.query and not p.fragment and re.search(r'відключ|знеструм|електроен',label,re.I) and not re.search(r'водопостач|води|квартир|садоч|екобус',label,re.I):urls.append(u)
  urls=list(dict.fromkeys(urls))[:4]
 seen=set()
 for url in urls:
  if time.monotonic()>deadline:report['stopReason']='time_budget';break
  try:
   raw=body if url==final else fetch(url)[0];page=BeautifulSoup(raw,'html.parser');published=published_at(page)
   if not published or not now-180*DAY<=published<=now:continue
   article=page.select_one('article .entry-content, article, main') or page
   report['articles'].append({'url':url,'publishedAt':published})
   for link in document_links(str(article),url,source['documentHosts']):
    if link['url'] in seen:continue
    if len(seen)>=3 or time.monotonic()>deadline:report['stopReason']='document_or_time_budget';break
    seen.add(link['url']);row={**link,'linkedFrom':url,'linkedFromPublishedAt':published,'documentPublishedAt':None,'status':'needs_document_review'}
    try:
     data,kind,doc_url=fetch(link['url'])
     if urllib.parse.urlsplit(doc_url).hostname not in source['documentHosts']:raise ValueError('Unreviewed redirect host')
     row.update(finalUrl=doc_url,sha256=hashlib.sha256(data).hexdigest(),bytes=len(data),**inspect_document(data,kind))
    except Exception as e:row.update(status='fetch_or_format_error',error=str(e)[:160])
    report['documents'].append(row)
  except Exception as e:report['errors'].append({'url':url,'error':str(e)[:160]})
 return report
